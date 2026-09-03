from __future__ import annotations

import csv
import hashlib
import io
import re
from html.parser import HTMLParser
from pathlib import Path

from .models import ParsedDocument, ParsedPage


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".html", ".htm", ".txt", ".md"}
PUBLISH_TIME_PATTERN = re.compile(
    r"(?:发布时间|发布日期|日期|publish(?:ed)?(?: time| date)?)\s*[：:]\s*"
    r"(\d{4}[-年/.]\d{1,2}[-月/.]\d{1,2}日?)",
    re.IGNORECASE,
)


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)

    def text(self) -> str:
        return "\n".join(self.parts)


class DocumentParser:
    def parse(self, path: str | Path) -> ParsedDocument:
        file_path = Path(path)
        if not file_path.exists():
            raise FileNotFoundError(file_path)
        suffix = file_path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {suffix}")

        doc_id = hashlib.sha256(file_path.read_bytes()).hexdigest()[:24]
        if suffix == ".pdf":
            pages = self._parse_pdf(file_path)
        elif suffix == ".docx":
            pages = self._parse_docx(file_path)
        elif suffix == ".xlsx":
            pages = self._parse_xlsx(file_path)
        elif suffix == ".csv":
            pages = [ParsedPage(self._parse_csv(file_path), None)]
        elif suffix in {".html", ".htm"}:
            pages = [ParsedPage(self._parse_html(file_path), None)]
        else:
            pages = [ParsedPage(file_path.read_text(encoding="utf-8"), None)]

        pages = [
            ParsedPage(self._sanitize_text(page.text), page.page_number)
            for page in pages
        ]

        joined = "\n".join(page.text for page in pages[:3])
        publish_match = PUBLISH_TIME_PATTERN.search(joined)
        publish_time = publish_match.group(1) if publish_match else None
        title = self._guess_title(file_path, joined)
        return ParsedDocument(
            doc_id=doc_id,
            title=title,
            source=file_path.name,
            publish_time=publish_time,
            pages=pages,
            metadata={"file_type": suffix.lstrip("."), "original_path": str(file_path)},
        )

    @staticmethod
    def _sanitize_text(text: str) -> str:
        """Remove NULs and replace malformed surrogate code points from PDF parsers."""
        return text.replace("\x00", "").encode("utf-8", errors="replace").decode("utf-8")

    @staticmethod
    def _guess_title(path: Path, text: str) -> str:
        for line in text.splitlines():
            candidate = line.strip().lstrip("#").strip()
            if 4 <= len(candidate) <= 100 and not candidate.startswith(("来源：", "日期：")):
                return candidate
        return path.stem

    @staticmethod
    def _parse_pdf(path: Path) -> list[ParsedPage]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support requires pypdf") from exc
        reader = PdfReader(str(path))
        return [
            ParsedPage((page.extract_text() or "").strip(), index + 1)
            for index, page in enumerate(reader.pages)
        ]

    @staticmethod
    def _parse_docx(path: Path) -> list[ParsedPage]:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support requires python-docx") from exc
        document = Document(str(path))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [ParsedPage("\n".join(parts), None)]

    @staticmethod
    def _parse_xlsx(path: Path) -> list[ParsedPage]:
        try:
            import pandas as pd
        except ImportError as exc:
            raise RuntimeError("XLSX support requires pandas and openpyxl") from exc
        sheets = pd.read_excel(path, sheet_name=None)
        pages: list[ParsedPage] = []
        for index, (sheet_name, frame) in enumerate(sheets.items(), start=1):
            pages.append(
                ParsedPage(
                    f"工作表：{sheet_name}\n{frame.fillna('').to_csv(index=False)}",
                    index,
                )
            )
        return pages

    @staticmethod
    def _parse_csv(path: Path) -> str:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)

    @staticmethod
    def _parse_html(path: Path) -> str:
        parser = _TextHTMLParser()
        parser.feed(path.read_text(encoding="utf-8"))
        return parser.text()


def parse_text(name: str, text: str, publish_time: str | None = None) -> ParsedDocument:
    """Convenience helper used by tests and integrations."""
    raw = text.encode("utf-8")
    doc_id = hashlib.sha256(raw).hexdigest()[:24]
    return ParsedDocument(
        doc_id=doc_id,
        title=name,
        source=name,
        publish_time=publish_time,
        pages=[ParsedPage(text=text, page_number=1)],
        metadata={"file_type": "text"},
    )
